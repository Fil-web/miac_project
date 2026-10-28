"""Формирование локальных выгрузок отчётов без внешних сервисов."""

from __future__ import annotations

from datetime import datetime
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from starlette.responses import Response


def _register_pdf_fonts() -> tuple[str, str]:
    """Подключает локальный шрифт с кириллицей из контейнера или macOS."""

    candidates = [
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
            Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ),
    ]
    for regular_path, bold_path in candidates:
        if regular_path.exists() and bold_path.exists():
            if "ReportRegular" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ReportRegular", str(regular_path)))
                pdfmetrics.registerFont(TTFont("ReportBold", str(bold_path)))
            return "ReportRegular", "ReportBold"
    raise RuntimeError("Не найден локальный шрифт с поддержкой кириллицы для PDF")


def _safe_text(value: object) -> str:
    """Преобразует значение отчёта в безопасную строку."""

    return "—" if value is None or value == "" else str(value)


def _xlsx_safe_text(value: object) -> str:
    """Не позволяет пользовательским строкам превратиться в формулы Excel."""

    text = _safe_text(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{text}"
    return text


def build_pdf_response(
    title: str,
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> Response:
    """Возвращает PDF с таблицей и встроенным кириллическим шрифтом."""

    regular_font, bold_font = _register_pdf_fonts()
    prepared_rows = [[_safe_text(value) for value in row] for row in rows]
    buffer = BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
        author="Учёт комплектующих АРМ МИАЦ",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=14,
        leading=17,
        alignment=TA_CENTER,
        spaceAfter=5 * mm,
    )
    cell_style = ParagraphStyle(
        "ReportCell",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=7.5,
        leading=9,
    )
    header_style = ParagraphStyle(
        "ReportHeader",
        parent=cell_style,
        fontName=bold_font,
        textColor=colors.white,
    )
    table_data = [
        [Paragraph(escape(_safe_text(value)), header_style) for value in headers],
        *[
            [Paragraph(escape(_safe_text(value)), cell_style) for value in row]
            for row in prepared_rows
        ],
    ]
    # Ширина столбцов распределяется по содержимому, но длинные поля ограничены.
    weights = []
    for index, header in enumerate(headers):
        values = [header, *(row[index] for row in prepared_rows if index < len(row))]
        weights.append(max(8, min(32, max(len(str(value)) for value in values))))
    available_width = page_size[0] - 20 * mm
    total_weight = sum(weights) or 1
    column_widths = [available_width * weight / total_weight for weight in weights]
    table = Table(table_data, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTNAME", (0, 1), (-1, -1), regular_font),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#adb5bd")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
            ]
        )
    )
    generated_at = datetime.now().strftime("%d.%m.%Y %H:%M")
    story = [
        Paragraph(escape(title), title_style),
        Paragraph(f"Сформировано: {generated_at}", cell_style),
        Spacer(1, 3 * mm),
        table,
    ]
    document.build(story)
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'},
    )


def build_xlsx_response(
    title: str,
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
) -> Response:
    """Возвращает оформленную книгу Excel с автофильтром и закреплением шапки."""

    prepared_rows = [[_xlsx_safe_text(value) for value in row] for row in rows]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"
    last_column = max(1, len(headers))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(bold=True, size=14, color="0D6EFD")
    title_cell.alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    sheet.append([])
    sheet.append(list(headers))
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6EFD")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in prepared_rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        values = [header, *(row[index - 1] for row in prepared_rows if index - 1 < len(row))]
        sheet.column_dimensions[get_column_letter(index)].width = max(
            12,
            min(45, max(len(str(value)) for value in values) + 2),
        )
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(last_column)}{max(4, 4 + len(prepared_rows))}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )
